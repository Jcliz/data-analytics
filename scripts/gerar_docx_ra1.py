"""Gera o documento do RA1 (.docx e, se houver LibreOffice, .pdf) a partir do template do professor (ADR-0003).

Uso: python scripts/gerar_docx_ra1.py [--template T.docx] [--fonte PASTA] [--saida S.docx] [--manter-notas] [--sem-pdf]
     Padrões: modelos/Template_RA1_Data_Analytics.docx, documentacao/ra1/, output/ra1-documentacao.docx
Precisa de python-docx (pip install python-docx).

Como o Markdown vira documento:
- 00-capa.md preenche os campos da capa; 01 a 07 preenchem as seções "RA1 / 01" a "RA1 / 07", na ordem.
- "# Título" identifica a seção; "## Subtítulo" casa com o subtítulo do template de mesmo nome.
- "**Campo:** valor" preenche o campo rotulado do template (ex.: "Unidade de observação").
- Tabelas Markdown preenchem a tabela do subtítulo, mantendo o cabeçalho e o estilo do template.
- Em "Registro das decisões", cada "### Dn" vira uma cópia do bloco de decisão do template; dentro dele,
  "#### Subtítulo" casa com os subtítulos do bloco.
- "- [x] texto" marca o item correspondente da "Conferência da entrega".
- Valor ainda entre colchetes ("[...]") deixa o campo do template como está; comentários <!-- --> são ignorados.
  Intervalos numéricos ("[0, 100]") contam como conteúdo. Links [texto](url) viram "texto (url)".
- Tabela Markdown sob um subtítulo que não tem tabela no template não entra no documento (o script avisa).
- As notas cinza de orientação do template são removidas (use --manter-notas para conferir o rascunho).
"""
import argparse
import copy
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

RAIZ = Path(__file__).resolve().parent.parent
CINZA = '606066'
REMOVIVEIS = {'anexos opcionais'}
INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*\s](?:[^*]*[^*\s])?\*)')
CAMPO = re.compile(r'^\*\*(?P<rotulo>[^*]+?):\*\*\s*(?P<valor>.*)$')
LINK = re.compile(r'\[([^\]\n]+)\]\((https?://[^)\s]+)\)')
INTERVALO = re.compile(r'^\[\s*-?[\d.,]+\s*[,;]\s*-?[\d.,]+\s*[\])]')


# ------------------------------------------------------------------ leitura do Markdown
def ler(caminho):
    texto = caminho.read_text(encoding='utf-8').replace('\r', '')
    texto = re.sub(r'<!--.*?-->', '', texto, flags=re.S)
    return LINK.sub(r'\1 (\2)', texto)


def estrutura(texto):
    """Devolve {'titulo', 'campos', 'subs': {nome: linhas}, 'grupos': {h2: [{titulo, campos, subs}]}}."""
    doc = {'titulo': None, 'campos': {}, 'subs': {}, 'grupos': {}, 'ordem': []}
    h2 = grupo = h4 = None
    for linha in texto.split('\n'):
        m = re.match(r'^(#{1,4})\s+(.*)$', linha)
        if m:
            nivel, nome = len(m.group(1)), m.group(2).strip()
            if nivel == 1:
                doc['titulo'] = nome
            elif nivel == 2:
                h2, grupo, h4 = nome, None, None
                doc['subs'].setdefault(h2, [])
                doc['ordem'].append(h2)
            elif nivel == 3:
                grupo = {'titulo': nome, 'campos': {}, 'subs': {}}
                doc['grupos'].setdefault(h2, []).append(grupo)
                h4 = None
            else:
                h4 = nome
                if grupo is not None:
                    grupo['subs'].setdefault(h4, [])
            continue
        campo = CAMPO.match(linha.strip())
        if grupo is not None:
            if campo:
                grupo['campos'][campo.group('rotulo').strip()] = campo.group('valor').strip()
            elif h4 is not None:
                grupo['subs'][h4].append(linha)
        else:
            if campo:
                doc['campos'][campo.group('rotulo').strip()] = campo.group('valor').strip()
            if h2 is not None:
                doc['subs'][h2].append(linha)
    return doc


def blocos(linhas, rotulos):
    """Converte linhas de Markdown em blocos ('p'|'li'|'ol'|'check', texto) e uma tabela opcional."""
    saida, tabela, paragrafo = [], None, []

    def fecha():
        if paragrafo:
            saida.append(('p', ' '.join(paragrafo)))
            paragrafo.clear()

    for linha in linhas:
        s = linha.strip()
        campo = CAMPO.match(s)
        if campo and campo.group('rotulo').strip() in rotulos:
            fecha()
            continue
        if not s:
            fecha()
        elif s.startswith('|'):
            fecha()
            celulas = [c.strip().replace('\\|', '|') for c in re.split(r'(?<!\\)\|', s.strip('|'))]
            if all(re.fullmatch(r':?-+:?', c) for c in celulas):
                continue
            tabela = (tabela or []) + [celulas]
        elif re.match(r'^- \[[ xX]\]\s', s):
            fecha()
            saida.append(('check', s))
        elif re.match(r'^[-*]\s', s):
            fecha()
            saida.append(('li', s[2:].strip()))
        elif re.match(r'^\d+\.\s', s):
            fecha()
            saida.append(('ol', re.sub(r'^\d+\.\s+', '', s)))
        else:
            paragrafo.append(s)
    fecha()
    corpo_tabela = tabela[1:] if tabela else None  # a primeira linha é o cabeçalho do Markdown
    return saida, corpo_tabela


def preenchido(valor):
    """Conteúdo de verdade: não vazio e não um espaço reservado "[...]" (intervalo numérico conta como conteúdo)."""
    v = (valor or '').lstrip()
    return bool(v) and (not v.startswith('[') or INTERVALO.match(v) is not None)


# ------------------------------------------------------------------ utilidades do .docx
def tag(el):
    return el.tag.split('}')[1]


def par(el, doc):
    return Paragraph(el, doc) if tag(el) == 'p' else None


def estilo(p):
    return p.style.name if p is not None and p.style is not None else ''


def cor(run):
    return str(run.font.color.rgb) if run.font.color is not None and run.font.color.type is not None else None


def e_nota(p):
    """Nota de orientação: texto cinza em itálico (ou o parágrafo de abertura da capa)."""
    if p is None or not p.text.strip():
        return False
    if p.text.startswith('Este documento acompanha a implementação'):
        return True
    return all(cor(r) == CINZA and r.italic for r in p.runs if r.text.strip())


def e_marcador(p):
    return p is not None and estilo(p) == 'Normal' and re.fullmatch(r'RA1\s+/\s+0\d', p.text.strip() or '') is not None


def rotulo(p):
    """Texto do rótulo de um campo ("Projeto  [..]"), ou None."""
    if p is None or len(p.runs) < 2 or not p.runs[0].bold:
        return None
    t = p.runs[0].text
    return t.strip() if t.endswith('  ') else None


def limpa_cor(rpr):
    if rpr is not None:
        for c in rpr.findall(qn('w:color')):
            rpr.remove(c)
        for i in rpr.findall(qn('w:i')):
            rpr.remove(i)
    return rpr


def escreve(paragrafo, texto, rpr_base=None):
    """Escreve texto com **negrito**, *itálico* e `código` no parágrafo."""
    for parte in filter(None, INLINE.split(texto)):
        if parte.startswith('**'):
            run = paragrafo.add_run(parte[2:-2]); run.bold = True
        elif parte.startswith('`'):
            run = paragrafo.add_run(parte[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(9)
        elif parte.startswith('*'):
            run = paragrafo.add_run(parte[1:-1]); run.italic = True
        else:
            run = paragrafo.add_run(parte)
        if rpr_base is not None:
            novo = copy.deepcopy(rpr_base)
            for filho in list(run._r.rPr if run._r.rPr is not None else []):
                if novo.find(filho.tag) is not None:
                    novo.remove(novo.find(filho.tag))
                novo.append(filho)
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, novo)


def novo_paragrafo(doc, depois_de, texto, estilo_nome='Normal'):
    p = doc.add_paragraph()
    p.style = doc.styles[estilo_nome]
    escreve(p, texto)
    depois_de.addnext(p._p)
    return p._p


# ------------------------------------------------------------------ preenchimento
class Gerador:
    def __init__(self, template):
        self.doc = Document(template)
        self.corpo = self.doc.element.body

    def elementos(self):
        return [el for el in self.corpo.iterchildren() if tag(el) in ('p', 'tbl')]

    def faixas(self):
        """Lista de faixas de elementos: [capa, seção 1, …, seção 7]."""
        els = self.elementos()
        cortes = [i for i, el in enumerate(els) if e_marcador(par(el, self.doc))]
        limites = [0] + cortes + [len(els)]
        return [els[a:b] for a, b in zip(limites, limites[1:])]

    # campos rotulados
    def preenche_campos(self, faixa, campos):
        for el in faixa:
            p = par(el, self.doc)
            r = rotulo(p)
            if r and r in campos and preenchido(campos[r]):
                rpr = limpa_cor(copy.deepcopy(p.runs[1]._r.rPr)) if p.runs[1]._r.rPr is not None else None
                for run in p.runs[1:]:
                    run._r.getparent().remove(run._r)
                escreve(p, campos[r], rpr)

    # corpo de um subtítulo: elementos até o próximo subtítulo, rótulo ou marcador
    def corpo_de(self, faixa, inicio):
        corpo = []
        for el in faixa[inicio + 1:]:
            p = par(el, self.doc)
            if p is not None and (estilo(p).startswith('Heading') or rotulo(p) or e_marcador(p)):
                break
            corpo.append(el)
        return corpo

    def preenche_tabela(self, tbl_el, linhas):
        tabela = Table(tbl_el, self.doc)
        # As linhas do template servem de modelo, em ciclo, para manter a alternância de cor das linhas
        modelos = [copy.deepcopy(r._tr) for r in tabela.rows[1:]] or [copy.deepcopy(tabela.rows[-1]._tr)]
        for r in tabela.rows[1:]:
            tbl_el.remove(r._tr)
        for k in range(len(linhas)):
            tbl_el.append(copy.deepcopy(modelos[k % len(modelos)]))
        for linha, row in zip(linhas, tabela.rows[1:]):
            for texto, cell in zip(linha, row.cells):
                p = cell.paragraphs[0]
                rpr = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
                for run in list(p.runs):
                    run._r.getparent().remove(run._r)
                for extra in cell.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
                escreve(p, texto, limpa_cor(rpr))

    def preenche_sub(self, faixa, nome, linhas, rotulos):
        alvo = next((i for i, el in enumerate(faixa)
                     if estilo(par(el, self.doc)).startswith('Heading 2')
                     and par(el, self.doc).text.strip().lower() == nome.strip().lower()), None)
        if alvo is None:
            print(f'  aviso: subtítulo "{nome}" não existe no template; ignorado')
            return
        corpo = self.corpo_de(faixa, alvo)
        conteudo, tabela = blocos(linhas, rotulos)
        if not any(preenchido(t) for _, t in conteudo) and not tabela:
            return
        tabelas = [el for el in corpo if tag(el) == 'tbl']
        if tabela and tabelas:
            self.preenche_tabela(tabelas[0], tabela)
        elif tabela:
            print(f'  aviso: "{nome}" não tem tabela no template; a tabela do Markdown não entrou (use texto ou lista)')
        checks = [b for b in conteudo if b[0] == 'check']
        caixas = [el for el in corpo if (par(el, self.doc) is not None and par(el, self.doc).text.startswith('☐'))]
        for (_, linha), el in zip(checks, caixas):
            if re.match(r'^- \[[xX]\]', linha):
                p = par(el, self.doc)
                p.runs[0].text = p.runs[0].text.replace('☐', '☒', 1)
        texto = [b for b in conteudo if b[0] != 'check']
        if texto:
            for el in corpo:  # remove os espaços reservados e as linhas em branco do template
                p = par(el, self.doc)
                quebra = el.findall('.//' + qn('w:br'))  # nunca apaga a quebra de página entre seções
                if p is not None and not quebra and (not p.text.strip() or p.text.strip().startswith('[')):
                    el.getparent().remove(el)
            ancora = faixa[alvo]
            for tipo, t in texto:
                nome_estilo = {'li': 'List Bullet', 'ol': 'List Number'}.get(tipo, 'Normal')
                ancora = novo_paragrafo(self.doc, ancora, t, nome_estilo)

    def remove_sub(self, faixa, nome):
        for i, el in enumerate(faixa):
            p = par(el, self.doc)
            if p is not None and estilo(p).startswith('Heading 2') and p.text.strip().lower() == nome:
                for alvo in [el] + self.corpo_de(faixa, i):
                    alvo.getparent().remove(alvo)
                return

    def preenche_grupos(self, faixa, grupos):
        """Duplica o bloco "Decisão e diagnóstico" do template, um por grupo (### Dn)."""
        inicio = next((i for i, el in enumerate(faixa) if rotulo(par(el, self.doc)) == 'Decisão e diagnóstico'), None)
        if inicio is None or not grupos:
            return
        fim = next(i for i, el in enumerate(faixa) if i > inicio and estilo(par(el, self.doc)) == 'Heading 2'
                   and par(el, self.doc).text.strip() == 'Resumo das intervenções')
        bloco = faixa[inicio:fim]
        resumo = faixa[fim]
        for grupo in grupos:
            copia = [copy.deepcopy(el) for el in bloco]
            for el in copia:  # cada cópia entra logo antes de "Resumo das intervenções", na ordem D1, D2…
                resumo.addprevious(el)
            rotulos = set(grupo['campos'])
            self.preenche_campos(copia, grupo['campos'])
            for nome, linhas in grupo['subs'].items():
                self.preenche_sub(self.atualiza(copia), nome, linhas, rotulos)
        for el in bloco:
            el.getparent().remove(el)

    def atualiza(self, faixa):
        """Recalcula a faixa depois de inserções (mesmo primeiro e último elemento ainda no corpo)."""
        els = self.elementos()
        vivos = [el for el in faixa if el.getparent() is not None]
        a, b = els.index(vivos[0]), els.index(vivos[-1])
        return els[a:b + 1]

    def capa(self, md):
        faixa = self.faixas()[0]
        self.preenche_campos(faixa, md['campos'])
        subtitulo = md['campos'].get('Subtítulo', 'Documentação da entrega')
        for el in faixa:
            p = par(el, self.doc)
            if estilo(p) == 'Subtitle' and p.runs:
                p.runs[0].text = subtitulo

    def secao(self, numero, md):
        faixa = self.faixas()[numero]
        h1 = next(par(el, self.doc).text.strip() for el in faixa if estilo(par(el, self.doc)) == 'Heading 1')
        if md['titulo'] and md['titulo'].strip().lower() != h1.lower():
            print(f'  aviso: título "{md["titulo"]}" difere do template ("{h1}")')
        rotulos = {rotulo(par(el, self.doc)) for el in faixa} - {None}
        self.preenche_campos(faixa, md['campos'])
        for h2, grupos in md['grupos'].items():
            self.preenche_grupos(faixa, grupos)
            faixa = self.faixas()[numero]
        for nome in md['ordem']:
            if nome in md['grupos']:
                continue
            self.preenche_sub(self.faixas()[numero], nome, md['subs'][nome], rotulos)
        presentes = {n.lower() for n in md['ordem']}
        for nome in REMOVIVEIS - presentes:
            self.remove_sub(self.faixas()[numero], nome)

    def quebras_de_pagina(self):
        """Troca o parágrafo de quebra de página antes de cada "RA1 / 0n" por "quebra antes" no próprio marcador.

        Assim, uma seção que ocupa a página inteira não deixa uma página em branco antes da seguinte.
        """
        for el in self.elementos():
            p = par(el, self.doc)
            if p is None or p.text.strip() or not el.findall('.//' + qn('w:br')):
                continue
            seguinte = el.getnext()
            if seguinte is not None and e_marcador(par(seguinte, self.doc)):
                par(seguinte, self.doc).paragraph_format.page_break_before = True
                el.getparent().remove(el)

    def remove_notas(self):
        for el in self.elementos():
            if e_nota(par(el, self.doc)):
                el.getparent().remove(el)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--template', type=Path, default=RAIZ / 'modelos' / 'Template_RA1_Data_Analytics.docx')
    ap.add_argument('--fonte', type=Path, default=RAIZ / 'documentacao' / 'ra1')
    ap.add_argument('--saida', type=Path, default=RAIZ / 'output' / 'ra1-documentacao.docx')
    ap.add_argument('--manter-notas', action='store_true', help='mantém as notas cinza de orientação do template')
    ap.add_argument('--sem-pdf', action='store_true', help='não tenta exportar o PDF com o LibreOffice')
    args = ap.parse_args()

    g = Gerador(args.template)
    g.quebras_de_pagina()  # antes de preencher, para nenhuma seção perder a página nova
    arquivos = sorted(args.fonte.glob('*.md'))
    capa = next((a for a in arquivos if a.name.startswith('00')), None)
    if capa:
        print(f'capa: {capa.name}')
        g.capa(estrutura(ler(capa)))
    for arquivo in arquivos:
        m = re.match(r'^0([1-7])-', arquivo.name)
        if m:
            print(f'seção {m.group(1)}: {arquivo.name}')
            g.secao(int(m.group(1)), estrutura(ler(arquivo)))
    if not args.manter_notas:
        g.remove_notas()

    args.saida.parent.mkdir(parents=True, exist_ok=True)
    g.doc.save(args.saida)
    print(f'gravado: {args.saida}')
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if soffice and not args.sem_pdf:
        subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', str(args.saida.parent), str(args.saida)],
                       check=False, capture_output=True, timeout=180)
        pdf = args.saida.with_suffix('.pdf')
        print(f'gravado: {pdf}' if pdf.exists() else 'aviso: o LibreOffice não gerou o PDF; exporte pelo Word')
    elif not args.sem_pdf:
        print('LibreOffice não encontrado: abra o .docx no Word e exporte em PDF')


if __name__ == '__main__':
    main()
