"""Monta os arquivos de entrega do RA1 e do RA2, em .docx e .pdf.

Uso: python scripts/montar_entregas.py [--resolucao DPI]
Precisa de python-docx, pypdf e do pdftoppm (pacote poppler-utils), que transforma as páginas do notebook em imagens.

- RA1: output/entrega-ra1.docx e output/entrega-ra1.pdf, só o documento do RA1 (output/ra1-documentacao.*).
- RA2: output/entrega-ra2.docx e output/entrega-ra2.pdf, o documento do RA1 seguido do notebook executado
  (output/projeto_data_analytics.pdf). No .pdf as páginas do notebook entram como estão; no .docx entram como
  imagens de página inteira, numa seção nova sem margens. Cada parte do .pdf vira um marcador no leitor de PDF.
Antes de montar, o script avisa se alguma parte está mais velha que a sua fonte.
"""
import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from pypdf import PdfReader, PdfWriter

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / 'output'
DOC_DOCX = SAIDA / 'ra1-documentacao.docx'
DOC_PDF = SAIDA / 'ra1-documentacao.pdf'
NOTEBOOK_PDF = SAIDA / 'projeto_data_analytics.pdf'
EQUIPE = 'Eduardo Fabri, João Correa, Rodrigo Alves e João Pedro Cardoso'  # CONTEXT.md, seção 1


def conferir_partes():
    partes = [
        (DOC_DOCX, sorted((RAIZ / 'documentacao' / 'ra1').glob('*.md')), 'python scripts/gerar_docx_ra1.py'),
        (DOC_PDF, [DOC_DOCX], 'python scripts/gerar_docx_ra1.py (ou exporte o .docx pelo Word)'),
        (NOTEBOOK_PDF, [RAIZ / 'notebook' / 'projeto_data_analytics.ipynb'],
         'exporte o notebook executado no Colab em PDF'),
    ]
    faltando = False
    for arquivo, fontes, como in partes:
        if not arquivo.exists():
            print(f'erro: falta {arquivo.relative_to(RAIZ)} ({como})')
            faltando = True
            continue
        for fonte in fontes:
            if fonte.exists() and fonte.stat().st_mtime > arquivo.stat().st_mtime:
                print(f'aviso: {fonte.relative_to(RAIZ)} mudou depois de {arquivo.name}; gere o arquivo de novo')
    if faltando:
        sys.exit(1)


def juntar_pdf(partes, saida, titulo):
    escritor = PdfWriter()
    for nome, pdf in partes:
        inicio = len(escritor.pages)
        for pagina in PdfReader(pdf).pages:
            escritor.add_page(pagina)
        escritor.add_outline_item(nome, inicio)
    escritor.add_metadata({'/Title': titulo, '/Author': EQUIPE})
    with open(saida, 'wb') as arquivo:
        escritor.write(arquivo)
    return len(escritor.pages)


def docx_com_notebook(saida, resolucao):
    documento = Document(DOC_DOCX)
    with tempfile.TemporaryDirectory() as pasta:
        subprocess.run(['pdftoppm', '-r', str(resolucao), '-png', str(NOTEBOOK_PDF), f'{pasta}/pagina'], check=True)
        imagens = sorted(Path(pasta).glob('pagina-*.png'))
        largura, altura = PdfReader(NOTEBOOK_PDF).pages[0].mediabox.width, PdfReader(NOTEBOOK_PDF).pages[0].mediabox.height
        secao = documento.add_section(WD_SECTION.NEW_PAGE)
        secao.page_width, secao.page_height = Mm(float(largura) * 25.4 / 72), Mm(float(altura) * 25.4 / 72)
        for margem in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin', 'header_distance',
                       'footer_distance', 'gutter'):
            setattr(secao, margem, Mm(0))
        for parte in (secao.header, secao.footer, secao.first_page_header, secao.first_page_footer):
            parte.is_linked_to_previous = False
            for paragrafo in parte.paragraphs:
                paragrafo.text = ''
        # add_section deixa um parágrafo vazio que só carrega a quebra da seção do documento; a quebra passa
        # para o último parágrafo do documento, para não sobrar uma página em branco antes do notebook.
        quebra = documento.paragraphs[-1]._p
        anterior = quebra.getprevious()
        if anterior is not None and anterior.tag == qn('w:p'):
            anterior.get_or_add_pPr().append(quebra.pPr.sectPr)
            quebra.getparent().remove(quebra)
        for i, imagem in enumerate(imagens):
            paragrafo = documento.add_paragraph()
            formato = paragrafo.paragraph_format
            formato.space_before = formato.space_after = Pt(0)
            formato.line_spacing = 1.0
            formato.page_break_before = i > 0
            paragrafo.add_run().add_picture(str(imagem), width=secao.page_width - Mm(1))
        documento.save(saida)
    return len(imagens)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--resolucao', type=int, default=150, help='DPI das páginas do notebook no .docx do RA2')
    args = ap.parse_args()
    if not shutil.which('pdftoppm'):
        sys.exit('erro: falta o pdftoppm (instale o poppler-utils)')
    conferir_partes()

    shutil.copyfile(DOC_DOCX, SAIDA / 'entrega-ra1.docx')
    paginas = juntar_pdf([('Documento do RA1', DOC_PDF)], SAIDA / 'entrega-ra1.pdf', 'Projeto Data Analytics — RA1')
    print(f'gravado: output/entrega-ra1.docx e output/entrega-ra1.pdf ({paginas} páginas)')

    paginas = juntar_pdf([('Documento do RA1', DOC_PDF), ('Notebook (RA1 e RA2)', NOTEBOOK_PDF)],
                         SAIDA / 'entrega-ra2.pdf', 'Projeto Data Analytics — RA1 e RA2')
    imagens = docx_com_notebook(SAIDA / 'entrega-ra2.docx', args.resolucao)
    print(f'gravado: output/entrega-ra2.pdf ({paginas} páginas) e output/entrega-ra2.docx '
          f'(documento + {imagens} páginas do notebook)')


if __name__ == '__main__':
    main()
